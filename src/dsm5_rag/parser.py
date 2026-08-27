"""文档解析器模块"""

import logging
from pathlib import Path
from typing import List

from llama_index.core import Document

logger = logging.getLogger(__name__)


class DocumentParser:
    """文档解析器 - 支持PDF、DOCX、TXT格式"""

    @staticmethod
    def parse_pdf(file_path: Path) -> List[Document]:
        """解析PDF文件"""
        try:
            import pypdf
        except ImportError:
            logger.error("请安装pypdf: pip install pypdf")
            return []

        documents = []
        try:
            with open(file_path, "rb") as file:
                pdf_reader = pypdf.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()

                    if text.strip():
                        doc = Document(
                            text=text,
                            metadata={
                                "file_name": file_path.name,
                                "file_type": "pdf",
                                "page": page_num + 1,
                                "total_pages": num_pages,
                                "source": str(file_path),
                            },
                        )
                        documents.append(doc)

            logger.info(
                f"PDF文件解析完成: {file_path.name}, 共{num_pages}页"
            )
            return documents

        except Exception as e:
            logger.error(f"PDF解析失败 {file_path}: {e}")
            return []

    @staticmethod
    def parse_docx(file_path: Path) -> List[Document]:
        """解析DOCX文件"""
        try:
            import docx
        except ImportError:
            logger.error("请安装python-docx: pip install python-docx")
            return []

        documents = []
        try:
            doc = docx.Document(file_path)
            full_text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text)
                    if row_text:
                        full_text.append(" | ".join(row_text))

            if full_text:
                combined_text = "\n".join(full_text)
                doc_obj = Document(
                    text=combined_text,
                    metadata={
                        "file_name": file_path.name,
                        "file_type": "docx",
                        "paragraphs": len(doc.paragraphs),
                        "tables": len(doc.tables),
                        "source": str(file_path),
                    },
                )
                documents.append(doc_obj)

            logger.info(
                f"DOCX文件解析完成: {file_path.name}, 生成{len(documents)}个文档"
            )
            return documents

        except Exception as e:
            logger.error(f"DOCX解析失败 {file_path}: {e}")
            return []

    @staticmethod
    def parse_txt(file_path: Path) -> List[Document]:
        """解析TXT文件"""
        documents = []
        try:
            with open(file_path, "r", encoding="utf-8") as file:
                content = file.read()

            if content.strip():
                doc = Document(
                    text=content,
                    metadata={
                        "file_name": file_path.name,
                        "file_type": "txt",
                        "source": str(file_path),
                    },
                )
                documents.append(doc)

            logger.info(
                f"TXT文件解析完成: {file_path.name}, 生成{len(documents)}个文档"
            )
            return documents

        except Exception as e:
            logger.error(f"TXT解析失败 {file_path}: {e}")
            return []